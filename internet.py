import os
import sys
import openai
import logging
import requests
from docx import Document
from PyPDF2 import PdfReader
from azure.storage.filedatalake import DataLakeServiceClient
from azure.ai.formrecognizer import DocumentAnalysisClient
from azure.core.credentials import AzureKeyCredential
import tempfile
import io
import re
import google.generativeai as genai
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# === Logging Setup ===
logging.basicConfig(
    level=logging.INFO,
    stream=sys.stdout,
    format="%(asctime)s - %(levelname)s - %(message)s"
)
logging.getLogger().addHandler(logging.StreamHandler(sys.stdout))

# === Environment Variables ===
SERP_API_KEY = os.environ.get("SERP_API_KEY")
GOOGLE_API_KEY = os.environ.get("GOOGLE_API_KEY")

genai.configure(api_key=GOOGLE_API_KEY)
GEMINI_MODEL = genai.GenerativeModel("models/gemini-1.5-pro-latest")


print("hello")

# === File Readers ===
def read_docx(file_path):
    try:
        doc = Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    except Exception as e:
        logging.error(f"[DOCX READ ERROR] {e}")
        return ""

def analyze_pdf_with_ai(pdf_bytes, filename="unknown.pdf"):
    try:
        endpoint = os.getenv("AZURE_FORM_RECOGNIZER_ENDPOINT")
        key = os.getenv("AZURE_FORM_RECOGNIZER_KEY")
        client = DocumentAnalysisClient(endpoint, AzureKeyCredential(key))
        poller = client.begin_analyze_document("prebuilt-document", document=pdf_bytes, content_type="application/pdf")
        result = poller.result()

        extracted_text = []
        for page in result.pages:
            for line in page.lines:
                extracted_text.append(line.content)
        for table in result.tables:
            extracted_text.append("\n--- Table ---")
            for cell in table.cells:
                extracted_text.append(f"Cell[{cell.row_index},{cell.column_index}]: {cell.content}")

        logging.info(f"[FORM RECOGNIZER] Extracted {len(extracted_text)} lines from {filename}")
        return "\n".join(extracted_text)
    except Exception as e:
        logging.error(f"[FORM RECOGNIZER ERROR] {filename} => {e}")
        return ""

def compute_cosine_similarity(query: str, documents: list[str], top_k: int = 3) -> list[str]:
    """Compute cosine similarity between the query and repository documents."""
    corpus = [query] + documents
    vectorizer = TfidfVectorizer().fit_transform(corpus)
    vectors = vectorizer.toarray()
    
    query_vector = vectors[0]
    doc_vectors = vectors[1:]
    
    similarity_scores = cosine_similarity([query_vector], doc_vectors)[0]
    top_indices = similarity_scores.argsort()[::-1][:top_k]
    
    return [documents[i] for i in top_indices]

def read_pdf(file_path):
    try:
        reader = PdfReader(file_path)
        return "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
    except Exception as e:
        logging.error(f"[PDF READ ERROR] {e}")
        return ""

# === Azure Data Lake Reader ===

def get_datalake_service_client():
    account_name = os.environ.get("AZURE_STORAGE_ACCOUNT_NAME")
    account_key = os.environ.get("AZURE_STORAGE_ACCOUNT_KEY")
    return DataLakeServiceClient(
        account_url=f"https://{account_name}.dfs.core.windows.net",
        credential=account_key
    )

def read_files_from_datalake():
    docs_info = []

    try:
        # === Load environment variables ===
        ACCOUNT_NAME = os.environ.get("AZURE_STORAGE_ACCOUNT_NAME")
        ACCOUNT_KEY = os.environ.get("AZURE_STORAGE_ACCOUNT_KEY")
        FILESYSTEM_NAME = os.environ.get("AZURE_DATA_LAKE_FILESYSTEM")

        if not all([ACCOUNT_NAME, ACCOUNT_KEY, FILESYSTEM_NAME]):
            logging.error("[DATALAKE] ❌ Missing one or more required environment variables.")
            return []

        # === Connect to Azure Data Lake ===
        logging.info(f"[DATALAKE] Connecting to: {ACCOUNT_NAME} | Filesystem: {FILESYSTEM_NAME}")
        service_client = DataLakeServiceClient(
            account_url=f"https://{ACCOUNT_NAME}.dfs.core.windows.net",
            credential=ACCOUNT_KEY
        )
        file_system_client = service_client.get_file_system_client(FILESYSTEM_NAME)

        # === Retrieve file paths ===
        logging.info("[DATALAKE] Fetching file paths...")
        paths = list(file_system_client.get_paths())
        logging.info(f"[DATALAKE] Found {len(paths)} paths.")

        for path in paths:
            logging.info(f"[DATALAKE] Path: {path.name} | Directory: {path.is_directory}")
            if path.is_directory:
                continue

            try:
                file_path = path.name
                file_client = file_system_client.get_file_client(file_path)
                download = file_client.download_file()
                file_data = download.readall()

                text = ""
                ext = file_path.lower().split('.')[-1]

                if ext == "pdf":
                    logging.info(f"[DATALAKE] Processing PDF: {file_path}")
                    text = analyze_pdf_with_ai(io.BytesIO(file_data), filename=file_path)

                elif ext == "docx":
                    logging.info(f"[DATALAKE] Processing DOCX: {file_path}")
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
                        tmp.write(file_data)
                        tmp.flush()
                        text = read_docx(tmp.name)

                else:
                    logging.info(f"[DATALAKE] Processing Text File: {file_path}")
                    try:
                        text = file_data.decode("utf-8", errors="ignore")
                    except Exception as decode_error:
                        logging.error(f"[DATALAKE DECODE ERROR] {file_path} => {decode_error}")
                        continue

                if text.strip():
                    char_count = len(text)
                    preview = text[:300].replace("\n", " ")
                    logging.info(f"[DATALAKE] ✅ Parsed: {file_path} ({char_count} chars)")
                    logging.info(f"[DATALAKE] Preview: {preview}...")
                    docs_info.append({"filename": file_path, "text": text})
                else:
                    logging.warning(f"[DATALAKE] ⚠️ No readable content in: {file_path}")

            except Exception as file_error:
                logging.error(f"[DATALAKE DOC READ ERROR] {path.name} => {file_error}")

    except Exception as e:
        logging.error(f"[DATALAKE CONNECTION ERROR] {e}")

    logging.info(f"[DATALAKE] Total files successfully processed: {len(docs_info)}")
    return docs_info

# === Helpers ===
# === Text Chunking ===
def chunk_text(text, max_words=4000):  # Leverage Gemini's high token limit
    words = text.split()
    return [' '.join(words[i:i + max_words]) for i in range(0, len(words), max_words)]

# === Summarization using Gemini ===
def summarize_text(text, max_tokens=800, model=None):
    try:
        model = model or genai.GenerativeModel("models/gemini-1.5-pro-latest")
        response = model.generate_content(f"Summarize this in {max_tokens} tokens:\n{text[:24000]}")
        return response.text.strip()
    except Exception as e:
        logging.error(f"[Gemini SUMMARY ERROR] {e}")
        return "Summary failed."


# === Web Search (unchanged) ===
def serpapi_search(query, max_results=5):
    try:
        params = {"engine": "google", "q": query, "api_key": SERP_API_KEY}
        response = requests.get("https://serpapi.com/search.json", params=params, timeout=10)
        return response.json().get("organic_results", [])[:max_results]
    except Exception as e:
        logging.error(f"[SERPAPI ERROR] {e}")
        return []

# === Word Count Utilities ===
def extract_requested_word_count(user_prompt):
    match = re.search(r"(\d{2,5})\s*words?", user_prompt.lower())
    return int(match.group(1)) if match else None

def calculate_max_tokens(word_count):
    return int(word_count * 1.5)

# === QA using Gemini ===
def answer_question_from_doc(document_text, user_question):
    chunks = chunk_text(document_text)

    for i, chunk in enumerate(chunks[:8]):
        try:
            logging.info(f"[QA Chunk {i+1}/{len(chunks)}] Searching for answer...")
            prompt = f"""
You are a helpful assistant. Answer the question strictly using the document content below.

Document:
\"\"\"{chunk}\"\"\"

Question:
{user_question}

If the answer is not found, say: "The answer is not available in the document."
"""
            model = genai.GenerativeModel("models/gemini-1.5-pro-latest")
            response = model.generate_content(prompt)
            answer = response.text.strip()
            if "not available" not in answer.lower():
                return answer
        except Exception as e:
            logging.error(f"[Gemini QA ERROR Chunk {i+1}] {e}")
    return "The answer is not available in the document."

# === Utility Functions ===
def limit_text_by_words(text, word_limit):
    words = text.split()
    return ' '.join(words[:word_limit])

def safe_concatenate_and_trim(docs, word_limit):
    combined = "\n\n".join(docs)
    return limit_text_by_words(combined, word_limit)

# === Proposal Generation ===
def generate_comprehensive_proposal(requirements_text, docs_info, user_prompt, use_internet, model=None):
    try:
        model = model or genai.GenerativeModel("models/gemini-1.5-pro-latest")

        # === Step 1: Summarize uploaded document ===
        summarized_requirements = summarize_text(requirements_text, 800, model=model)
        uploaded_doc_text = safe_concatenate_and_trim(
            [doc.get("text", "") for doc in docs_info if doc.get("text")],
            word_limit=3000
        )
        logging.info(f"[UPLOAD] Uploaded document content preview: {uploaded_doc_text[:300].replace(chr(10), ' ')}...")

        # === Step 2: Classify prompt intent ===
        logging.info("[STEP 2] Classifying user prompt intent")
        intent_prompt = f"""
Classify this prompt into one of the following:
- question-about-uploaded-document
- solution-needed-from-repo
- full-context

Prompt:
{user_prompt.strip()}
"""
        intent_response = model.generate_content(intent_prompt)
        mode = intent_response.text.strip().lower()
        logging.info(f"[INTENT] Classified user prompt as: {mode}")

        # === Step 3: Gather repository content with cosine similarity ===
        logging.info("[STEP 3] Reading and matching Azure Data Lake documents")
        azure_context = ""
        if mode in ["repository-needed", "solution-needed", "full-context"]:
            azure_docs = read_files_from_datalake()
            logging.info(f"[REPO] Total documents read from Data Lake: {len(azure_docs)}")

            repo_texts = [doc.get("text", "") for doc in azure_docs if doc.get("text")]
            top_matches = compute_cosine_similarity(summarized_requirements, repo_texts, top_k=3)
            if top_matches:
                azure_context = safe_concatenate_and_trim(top_matches, word_limit=3000)
                logging.info(f"[REPO] Selected top {len(top_matches)} documents using cosine similarity.")
            else:
                azure_context = "No strong repository content match found."
                logging.warning("[REPO] No matches found using cosine similarity.")

        # === Step 4: Gather internet data ===
        logging.info("[STEP 4] Searching internet context (if required)")
        internet_data = ""
        if use_internet and mode == "full-context":
            serp_results = serpapi_search(summarized_requirements)
            for result in serp_results:
                snippet_summary = summarize_text(f"{result.get('title')} - {result.get('snippet')}", 150)
                internet_data += (
                    f"Source: {result.get('link')}\n"
                    f"Title: {result.get('title')}\n"
                    f"Snippet: {result.get('snippet')}\n"
                    f"Summary: {snippet_summary}\n\n"
                )

        # === Step 5: Calculate max token allowance ===
        logging.info("[STEP 5] Calculating max token allowance")
        requested_words = extract_requested_word_count(user_prompt)
        dynamic_max_tokens = min(calculate_max_tokens(requested_words), 7000) if requested_words else 3500

        word_instruction = ""
        if requested_words:
            word_instruction = (
                f"IMPORTANT: Your response must be at least {requested_words} words. "
                "Do not stop early. Expand fully until reaching the requested word count."
            )

        # === Step 6: Build final prompt ===
        logging.info("[STEP 6] Constructing final prompt for Gemini")
        sections = [
            word_instruction,
            f"# User Prompt\n{user_prompt.strip()}",
            f"# Requirements Summary\n{summarized_requirements.strip()}",
            f"# Uploaded Document Context\n{uploaded_doc_text.strip()}"
        ]
        if azure_context:
            sections.append(f"# Repository Insights\n{azure_context.strip()}")
        if internet_data:
            sections.append(f"# Internet Findings\n{internet_data.strip()}")

        final_prompt = "\n\n".join(sections)
        logging.info(f"[GEMINI] Prompt word count: {len(final_prompt.split())}, max_tokens (approx): {dynamic_max_tokens}")

        # === Step 7: Call Gemini ===
        logging.info("[STEP 7] Calling Gemini to generate final proposal")
        response = model.generate_content(final_prompt)
        return response.text.strip()

    except Exception as e:
        logging.error(f"[generate_comprehensive_proposal ERROR] {e}")
        return "Unable to generate a response due to an internal error."

